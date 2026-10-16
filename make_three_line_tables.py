"""
Generate publication-style three-line tables (三线表) for Tables 2, 3 and 4 of
Card & Krueger (1994), from the already-produced `table*_updated.csv` files.

For each table the script writes two files into output/three_line_tables/:

  * a PNG image (matplotlib, crisp vector-like rendering at 200 dpi), and
  * an editable .docx (python-docx) with a real Word three-line table
    (top rule, header rule, bottom rule; no vertical or internal rules).

The source CSVs (output/table*_updated.csv) are only read, never modified.
"""

from __future__ import annotations

import os
import textwrap

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "output")
DEST = os.path.join(OUT, "three_line_tables")
os.makedirs(DEST, exist_ok=True)

EM = "—"  # em dash used for missing cells

# ---------------------------------------------------------------------------
# Number formatting helpers
# ---------------------------------------------------------------------------
def _fmt(v: float, nd: int) -> str:
    if pd.isna(v):
        return EM
    return f"{v:.{nd}f}"


def _se(se: float, nd: int) -> str:
    """Standard error in parentheses; a zero SE (degenerate group) is a dash."""
    if pd.isna(se) or se == 0.0:
        return EM
    return f"({se:.{nd}f})"


def _stack(mean: float, se: float, nd_mean: int, nd_se: int) -> str:
    """Mean with standard error stacked underneath (the paper's cell layout)."""
    if pd.isna(se) or se == 0.0:
        return f"{_fmt(mean, nd_mean)}\n{EM}"
    return f"{_fmt(mean, nd_mean)}\n({se:.{nd_se}f})"


# ---------------------------------------------------------------------------
# Table specs  (cell = dict(text, span, align, bold, italic))
# ---------------------------------------------------------------------------
def _cell(text: str, span: int = 1, align: str = "center",
          bold: bool = False, italic: bool = False) -> dict:
    return {"text": text, "span": span, "align": align,
            "bold": bold, "italic": italic}


def build_table2() -> dict:
    df = pd.read_csv(os.path.join(OUT, "table2_updated.csv"))

    dec = {  # variable -> (mean dp, se dp)
        "Burger King": (1, 1), "KFC": (1, 1), "Roy Rogers": (1, 1),
        "Wendy's": (1, 1), "Company-owned": (1, 1),
        "FTE employment": (2, 2),
        "Percentage full-time employees": (1, 1),
        "Starting wage": (2, 2),
        "Wage = $4.25 (percentage)": (1, 1),
        "Price of full meal": (2, 2),
        "Hours open (weekday)": (2, 2),
        "Recruiting bonus": (1, 1),
        "Wage = $5.05 (percentage)": (1, 1),
    }

    sections = [
        ("1. Distribution of store types (percentages):",
         ["Burger King", "KFC", "Roy Rogers", "Wendy's", "Company-owned"]),
        ("2. Means in wave 1:",
         ["FTE employment", "Percentage full-time employees", "Starting wage",
          "Wage = $4.25 (percentage)", "Price of full meal",
          "Hours open (weekday)", "Recruiting bonus"]),
        ("3. Means in wave 2:",
         ["FTE employment", "Percentage full-time employees", "Starting wage",
          "Wage = $4.25 (percentage)", "Wage = $5.05 (percentage)",
          "Price of full meal", "Hours open (weekday)", "Recruiting bonus"]),
    ]

    body: list[list[dict]] = []
    letter = 0
    for sec_title, vars_ in sections:
        body.append([_cell(sec_title, span=4, align="left", italic=True)])
        for v in vars_:
            r = df[df["Variable"] == v].iloc[0]
            nd_m, nd_s = dec[v]
            letter += 1
            # t statistic: dash when either group has zero variance
            t = r["t"]
            t_txt = EM if (r["NJ_se"] == 0.0 or r["PA_se"] == 0.0) else f"{t:.1f}"
            body.append([
                _cell(f"{chr(96 + letter)}. {v}", align="left"),
                _cell(_stack(r["NJ"], r["NJ_se"], nd_m, nd_s)),
                _cell(_stack(r["PA"], r["PA_se"], nd_m, nd_s)),
                _cell(t_txt),
            ])

    return {
        "title": "TABLE 2—MEANS OF SELECTED OUTCOME VARIABLES",
        "header": [
            [_cell("Variable", align="left"), _cell("Stores in:", span=2),
             _cell("t^a")],
            [_cell("", align="left"), _cell("NJ"), _cell("PA"), _cell("")],
        ],
        "body": body,
        "col_widths": [0.40, 0.21, 0.21, 0.18],
        "notes": [
            "Notes: Standard errors are given in parentheses. See text for "
            "definitions of the variables.",
            "^a Test of equality of means in New Jersey and Pennsylvania.",
        ],
    }


def build_table3() -> dict:
    df = pd.read_csv(os.path.join(OUT, "table3_updated.csv"))
    nd = 2  # FTE employment, two decimals throughout

    def vals(row_label: str) -> list[dict]:
        r = df[df["Row"] == row_label].iloc[0]
        return [
            _cell(_stack(r["PA"], r["PA_se"], nd, nd)),
            _cell(_stack(r["NJ"], r["NJ_se"], nd, nd)),
            _cell(_stack(r["NJ-PA"], r["NJ-PA_se"], nd, nd)),
            _cell(_stack(r["Low ($4.25)"], r["Low ($4.25)_se"], nd, nd)),
            _cell(_stack(r["Mid ($4.26-$4.99)"], r["Mid ($4.26-$4.99)_se"], nd, nd)),
            _cell(_stack(r["High ($5.00+)"], r["High ($5.00+)_se"], nd, nd)),
            _cell(_stack(r["Low-High"], r["Low-High_se"], nd, nd)),
            _cell(_stack(r["Mid-High"], r["Mid-High_se"], nd, nd)),
        ]

    body = [
        [_cell("1. FTE employment before, all available observations", align="left")] + vals(
            "1. FTE employment before, all available observations"),
        [_cell("2. FTE employment after, all available observations", align="left")] + vals(
            "2. FTE employment after, all available observations"),
        [_cell("3. Change in mean FTE employment", align="left")] + vals(
            "3. Change in mean FTE employment"),
        [_cell("4. Change in mean FTE employment, balanced sample of stores^c", align="left")] + vals(
            "4. Change in mean FTE employment, balanced sample"),
        [_cell("5. Change in mean FTE employment, setting FTE at temporarily "
               "closed stores to 0^d", align="left")] + vals(
            "5. Change, temporarily closed stores set to 0"),
    ]

    return {
        "title": "TABLE 3—AVERAGE EMPLOYMENT PER STORE BEFORE AND AFTER "
                 "THE RISE IN NEW JERSEY MINIMUM WAGE",
        "header": [
            [_cell("Variable", align="left"), _cell("Stores by state", span=3),
             _cell("Stores in New Jersey^a", span=3),
             _cell("Differences within NJ^b", span=2)],
            [_cell("", align="left"), _cell("PA (i)"), _cell("NJ (ii)"),
             _cell("Difference,\nNJ-PA (iii)"), _cell("Wage =\n$4.25 (iv)"),
             _cell("Wage =\n$4.26-$4.99 (v)"), _cell("Wage ≥\n$5.00 (vi)"),
             _cell("Low-high\n(vii)"), _cell("Midrange-\nhigh (viii)")],
        ],
        "body": body,
        "col_widths": [0.28, 0.09, 0.09, 0.10, 0.09, 0.10, 0.09, 0.08, 0.08],
        "notes": [
            "Notes: Standard errors are shown in parentheses. The sample consists "
            "of all stores with available data on employment. FTE (full-time-"
            "equivalent) employment counts each part-time worker as half a "
            "full-time worker. Employment at six closed stores is set to zero; "
            "employment at four temporarily closed stores is treated as missing.",
            "^a Stores in New Jersey are classified by starting wage in wave 1: "
            "$4.25 (N = 101), $4.26-$4.99 (N = 140), or ≥ $5.00 (N = 73).",
            "^b Difference in employment between low-wage ($4.25) and high-wage "
            "(≥ $5.00) stores, and between midrange ($4.26-$4.99) and "
            "high-wage stores.",
            "^c Subset of stores with available employment data in waves 1 and 2.",
            "^d In this row only, wave-2 employment at four temporarily closed "
            "stores is set to 0; changes are based on stores with employment data "
            "in both waves.",
        ],
    }


def build_table4() -> dict:
    df = pd.read_csv(os.path.join(OUT, "table4_updated.csv"))
    m = {row["Model"]: row for _, row in df.iterrows()}

    def coef(model: str) -> dict:
        r = m[model]
        return _cell(_stack(r["NJ"], r["NJ_se"], 2, 2))

    def gap(model: str) -> dict:
        r = m[model]
        return _cell(_stack(r["gap"], r["gap_se"], 2, 2))

    def pval(model: str) -> dict:
        r = m[model]
        return _cell(EM if pd.isna(r["p_controls"]) else f"{r['p_controls']:.2f}")

    def ser(model: str) -> dict:
        return _cell(f"{m[model]['SER']:.2f}")

    body = [
        [_cell("1. New Jersey dummy", align="left"),
         coef("(i)"), coef("(ii)"), _cell(EM), _cell(EM), _cell(EM)],
        [_cell("2. Initial wage gap^a", align="left"),
         _cell(EM), _cell(EM), gap("(iii)"), gap("(iv)"), gap("(v)")],
        [_cell("3. Controls for chain and ownership^b", align="left"),
         _cell("no"), _cell("yes"), _cell("no"), _cell("yes"), _cell("yes")],
        [_cell("4. Controls for region^c", align="left"),
         _cell("no"), _cell("no"), _cell("no"), _cell("no"), _cell("yes")],
        [_cell("5. Standard error of regression", align="left"),
         ser("(i)"), ser("(ii)"), ser("(iii)"), ser("(iv)"), ser("(v)")],
        [_cell("6. Probability value for controls^d", align="left"),
         pval("(i)"), pval("(ii)"), pval("(iii)"), pval("(iv)"), pval("(v)")],
    ]

    return {
        "title": "TABLE 4—REDUCED-FORM MODELS FOR CHANGE IN EMPLOYMENT",
        "header": [
            [_cell("Independent variable", align="left"), _cell("(i)"),
             _cell("(ii)"), _cell("(iii)"), _cell("(iv)"), _cell("(v)")],
        ],
        "body": body,
        "col_widths": [0.40, 0.12, 0.12, 0.12, 0.12, 0.12],
        "notes": [
            "Notes: Standard errors are given in parentheses. The sample consists "
            "of 357 stores with available data on employment and starting wages "
            "in waves 1 and 2. The dependent variable in all models is change in "
            "FTE employment (mean −0.238, standard deviation 8.825). All "
            "models include an unrestricted constant (not reported).",
            "^a Proportional increase in starting wage required to reach the new "
            "minimum wage; the wage gap is 0 for Pennsylvania stores.",
            "^b Three dummy variables for chain type and a dummy for company "
            "ownership.",
            "^c Dummy variables for three regions of New Jersey (North, Central, "
            "South) and two regions of eastern Pennsylvania.",
            "^d Probability value of the joint F test for exclusion of all control "
            "variables.",
        ],
    }


# ---------------------------------------------------------------------------
# PNG rendering (manual: text + three horizontal rules)
# ---------------------------------------------------------------------------
def render_png(spec: dict, path: str, figsize, fs_title, fs_header, fs_body,
               fs_notes) -> None:
    title = spec["title"]
    header = spec["header"]
    body = spec["body"]
    colw = spec["col_widths"]
    notes = spec.get("notes", [])
    n_cols = len(colw)

    # Row heights (arbitrary units).
    H_HDR, H_2, H_1 = 1.0, 1.0, 0.72
    heights = [H_HDR] * len(header)
    for r in body:
        multi = any("\n" in c["text"] for c in r)
        heights.append(H_2 if multi else H_1)
    total_h = sum(heights)

    # Wrap title and notes so nothing runs off the figure.
    title_lines = textwrap.wrap(title, width=95) or [title]
    note_lines = []
    for n in notes:
        note_lines.append(textwrap.wrap(n, width=135) or [n])

    fig, ax = plt.subplots(figsize=figsize)
    ax.axis("off")
    ax.set_xlim(0, 1)

    # Reserve vertical space: title, table, notes.
    gap = 0.018
    title_space = 0.012 + 0.030 * len(title_lines)
    note_space = 0.010 + 0.024 * sum(len(ls) for ls in note_lines)
    table_top = 1.0 - title_space
    table_bottom = note_space
    scale = (table_top - table_bottom) / total_h

    # Title.
    y = 1.0 - 0.010
    for ln in title_lines:
        ax.text(0.5, y, ln, ha="center", va="top", fontsize=fs_title,
                fontweight="bold")
        y -= 0.030

    # Column x boundaries.
    xb = [0.0]
    for w in colw:
        xb.append(xb[-1] + w)

    # Rows.
    y = table_top
    for ri, row in enumerate(header + body):
        h = heights[ri]
        yb = y - h * scale
        is_header = ri < len(header)
        fs = fs_header if is_header else fs_body
        cx = 0
        for cell in row:
            span = cell.get("span", 1)
            x0, x1 = xb[cx], xb[cx + span]
            cx += span
            align = cell.get("align", "center")
            bold = cell.get("bold", False)
            italic = cell.get("italic", False)
            if align == "left":
                x, ha = x0 + 0.012, "left"
            elif align == "right":
                x, ha = x1 - 0.012, "right"
            else:
                x, ha = (x0 + x1) / 2, "center"
            ax.text(x, (y + yb) / 2, cell["text"], ha=ha, va="center",
                    fontsize=fs, fontweight="bold" if bold else "normal",
                    fontstyle="italic" if italic else "normal")
        y = yb

    # The three rules.
    top_y = table_top
    hdr_y = table_top - sum(heights[:len(header)]) * scale
    bot_y = table_top - total_h * scale
    ax.plot([0, 1], [top_y, top_y], color="black", lw=1.5, clip_on=False)
    ax.plot([0, 1], [hdr_y, hdr_y], color="black", lw=0.8, clip_on=False)
    ax.plot([0, 1], [bot_y, bot_y], color="black", lw=1.5, clip_on=False)

    # Notes under the bottom rule.
    y = bot_y - 0.022
    for ls in note_lines:
        for ln in ls:
            ax.text(0.0, y, ln, ha="left", va="top", fontsize=fs_notes)
            y -= 0.024

    fig.savefig(path, dpi=200, bbox_inches="tight", facecolor="white")
    plt.close(fig)


# ---------------------------------------------------------------------------
# DOCX rendering (python-docx three-line table)
# ---------------------------------------------------------------------------
def _set_cell_border(cell, top=None, bottom=None, left=None, right=None) -> None:
    """Set explicit borders on one cell (None = no border)."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge, spec in (("top", top), ("bottom", bottom),
                       ("left", left), ("right", right)):
        el = tcBorders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tcBorders.append(el)
        if spec is None:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(spec))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "000000")


THICK = 12   # 1.5 pt (eighths of a point)
THIN = 6     # 0.75 pt


def _write_cell(cell, text: str, align: str, bold: bool, italic: bool,
                size: float) -> None:
    cell.text = ""
    para = cell.paragraphs[0]
    para.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT,
                      "center": WD_ALIGN_PARAGRAPH.CENTER,
                      "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    para.paragraph_format.space_before = Pt(1)
    para.paragraph_format.space_after = Pt(1)
    lines = text.split("\n")
    for i, ln in enumerate(lines):
        run = para.add_run(ln)
        run.bold = bold
        run.italic = italic
        run.font.size = Pt(size)
        if i < len(lines) - 1:
            run.add_break()


def render_doc(spec: dict, path: str, landscape: bool = False) -> None:
    doc = Document()

    # Page orientation.
    sec = doc.sections[0]
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        w, h = sec.page_width, sec.page_height
        sec.page_width, sec.page_height = h, w

    # Title.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(spec["title"])
    run.bold = True
    run.font.size = Pt(11)

    header = spec["header"]
    body = spec["body"]
    n_cols = len(spec["col_widths"])
    n_header = len(header)

    table = doc.add_table(rows=n_header + len(body), cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Fill cells (merge horizontal spans).
    for ri, row in enumerate(header + body):
        ci = 0
        for cell_spec in row:
            span = cell_spec.get("span", 1)
            cell = table.cell(ri, ci)
            if span > 1:
                cell = cell.merge(table.cell(ri, ci + span - 1))
            _write_cell(cell, cell_spec["text"], cell_spec.get("align", "center"),
                        cell_spec.get("bold", False),
                        cell_spec.get("italic", False), 8.5)
            ci += span

    # Three-line borders.
    last = len(table.rows) - 1
    for ri, row in enumerate(table.rows):
        for cell in row.cells:
            top = THICK if ri == 0 else None
            bottom = None
            if ri == n_header - 1:
                bottom = THIN
            if ri == last:
                bottom = THICK
            _set_cell_border(cell, top=top, bottom=bottom, left=None, right=None)

    # Column widths (approximate, in inches).
    total_w = Inches(9.0) if landscape else Inches(6.5)
    widths = [w * total_w for w in spec["col_widths"]]
    for ci, w in enumerate(widths):
        for ri in range(len(table.rows)):
            try:
                table.cell(ri, ci).width = w
            except Exception:
                pass

    # Notes.
    for n in spec.get("notes", []):
        p = doc.add_paragraph()
        run = p.add_run(n)
        run.font.size = Pt(8)

    doc.save(path)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main() -> None:
    jobs = [
        (build_table2(), "table2_three_line", (9.6, 11.0), 11, 8.5, 8.0, 7.0,
         False),
        (build_table3(), "table3_three_line", (14.0, 5.8), 10.5, 8.0, 7.6, 7.0,
         True),
        (build_table4(), "table4_three_line", (9.6, 4.6), 11, 8.5, 8.5, 7.0,
         False),
    ]
    for spec, name, figsize, fst, fsh, fsb, fsn, landscape in jobs:
        png_path = os.path.join(DEST, f"{name}.png")
        doc_path = os.path.join(DEST, f"{name}.docx")
        render_png(spec, png_path, figsize, fst, fsh, fsb, fsn)
        render_doc(spec, doc_path, landscape=landscape)
        print(f"[ok] {name}.png + {name}.docx")


if __name__ == "__main__":
    main()
